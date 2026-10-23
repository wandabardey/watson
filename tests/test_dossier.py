"""Proposition écrite : parties rédigées, tableaux de la base (objets, cycle de vie, détail CRUD, maquettes, chiffrage),
proposition de Watson, exports Word, HTML et Markdown."""

import json
import shutil
from pathlib import Path

import pytest

from referentiel import chiffrage as ch
from referentiel import dossier as do

FRET = Path(__file__).resolve().parent.parent / "exemples" / "fret"
REPONSE = {"titre": "Fret \u2014 Proposition 4SH", "redige": {"besoin": "Le client veut **suivre** ses expéditions.\n\n| Niveau | Contenu |\n|---|---|\n| Socle | Tout |",
                                                        "dispositif": "1. **Cadrage** : ateliers.", "inconnu": "x"},
           "fonctionnalites": [{"id": "F1", "nom": "Saisie des expéditions", "utilisateurs": "Expéditeur", "description": "Saisir un envoi.", "lot": "MVP",
                                "scenarios": ["SC-INCONNU"], "ecrans": ["SCR-001", "SCR-999"]}, {"nom": "Sans id"}, {"id": "F9"}],
           "points": [{"theme": "Budget", "question": "Une enveloppe existe-t-elle ?"}, {"theme": "x"}], "a_confirmer": ["Volumétrie"]}


@pytest.fixture()
def projet(tmp_path, monkeypatch):
    monkeypatch.setenv("WATSON_DONNEES", str(tmp_path / "donnees"))
    monkeypatch.setenv("HOME", str(tmp_path / "maison"))
    (tmp_path / "maison" / "Downloads").mkdir(parents=True)
    racine = tmp_path / "fret"
    shutil.copytree(FRET, racine, ignore=shutil.ignore_patterns("sortie", "recette"))
    return racine


def test_document_sans_redaction(projet):
    md = do.document(projet)["markdown"]
    assert "## 1. Récap du besoin" in md and "*Récap du besoin : à rédiger.*" in md
    assert "| Expédition | Saisie dans l'application |" in md and "### 2.3 Cycle de vie : expédition" in md
    assert "| Expédition | Saisie d'une expédition" in md and "C : l'expéditeur" in md
    assert "*Chiffrage à faire dans l'onglet Chiffrage.*" in md and "## 6. Points ouverts" in md


def test_rediger_garder_et_exporter(projet):
    sdk = pytest.importorskip("claude_agent_sdk")
    vus = []

    async def requete(prompt, options):
        vus.append(options.system_prompt)
        yield sdk.ResultMessage(subtype="success", duration_ms=1, duration_api_ms=1, is_error=False, num_turns=1, session_id="s", result=json.dumps(REPONSE))
    ch.enregistrer(projet, {"tjm": 640, "profils": [{"nom": "PO", "jours_par_sprint": 5}], "phases": [{"id": "mvp", "nom": "MVP", "modalite": "sprints", "sprints": 4}],
                            "annonce": ["mvp"], "options": ["Extraction automatisée"], "jalons": [{"semaine": 4, "libelle": "Démo du MVP"}]})
    v = do.rediger("p1", projet, requete=requete)
    assert "Watson y insère lui-même les tableaux" in vus[0] and "MVP : Sprints, 4 sprints, 12 800 € HT" in vus[0]
    att = v["en_attente"]["dossier"]
    assert set(att["redige"]) == {"besoin", "dispositif"} and [f["id"] for f in att["fonctionnalites"]] == ["F1", "F2"]
    assert att["fonctionnalites"][0]["ecrans"] == ["SCR-001"] and att["fonctionnalites"][0]["scenarios"] == []
    assert len(att["points"]) == 1 and "\u2014" not in att["titre"] and "suivre" in v["apercu_attente"]
    v = do.garder("p1", projet, "Wanda")
    do.enregistrer(projet, do.lire(projet))
    assert v["en_attente"] is None and v["dossier"]["auteur"] == "Wanda"
    md = v["markdown"]
    assert "| F1 | Saisie des expéditions | Expéditeur | Saisir un envoi. | MVP |" in md
    assert "| MVP | Sprints | 4 sprints | 12 800 € HT (20 j) |" in md and "Extraction automatisée" in md
    assert "**Première enveloppe à annoncer :** MVP, soit 12 800 € HT." in md and "- Semaine 4 : Démo du MVP" in md
    assert "- ☐ **Budget :** Une enveloppe existe-t-elle ?" in md
    h = Path(do.exporter(projet, "html")["chemin"]).read_text(encoding="utf-8")
    assert "<table>" in h and "<strong>suivre</strong>" in h and "<th>Niveau</th>" in h
    w = do.exporter(projet, "docx")
    from docx import Document
    doc = Document(w["chemin"])
    assert any("Récap du besoin" in p.text for p in doc.paragraphs) and len(doc.tables) >= 4
    assert do.exporter(projet, "md")["nom"].endswith(".md")
